"""큐텐 대량 등록 엑셀 양식 자동 채우기.

원본: Qoo10_EditItemList.xlsx (50 컬럼)
행 구조:
  1행: 영문 필드명 (item_number, price_yen 등)
  2행: 한글 필드명
  3행: 필수/선택 입력 여부
  4행: 각 필드 설명
  5행~: 실제 데이터 입력 행
"""
from __future__ import annotations

import io
from copy import copy
from datetime import date, timedelta
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

_TEMPLATE_PATH = Path(__file__).resolve().parent.parent / "data" / "qoo10_template.xlsx"

DATA_START_ROW = 5  # 템플릿 구조상 5행부터 실제 상품 데이터


# 엑셀 컬럼 인덱스 (1-based). 필요한 것만.
COL = {
    "item_number": 1,          # A
    "seller_unique_item_id": 2, # B
    "category_number": 3,      # C
    "brand_number": 4,         # D
    "item_name": 5,            # E
    "item_status": 7,          # G
    "end_date": 9,             # I
    "price_yen": 10,           # J
    "quantity": 13,            # M
    "image_main_url": 17,      # Q
    "item_description": 24,    # X
    "shipping_number": 25,     # Y
    "available_shipping_date": 27, # AA
    "search_keyword": 29,      # AC
    "item_condition_type": 30, # AD
    "origin_type": 31,         # AE
    "origin_country_id": 33,   # AG
    "item_weight": 36,         # AJ
    "under18s_display": 46,    # AT
}


def _set(ws, row: int, key: str, value: Any) -> None:
    if value is None or value == "":
        return
    col = COL.get(key)
    if col:
        ws.cell(row=row, column=col, value=value)


def build_qoo10_excel(
    rows: list[dict],
    defaults: dict,
) -> bytes:
    """시트 행 + 기본값을 받아 큐텐 대량등록 xlsx를 생성.

    rows: [{ product_name, sell_price_jpy, cover_image_url, weight_g, ... }]
    defaults: {
        category_number, brand_number, shipping_number,
        end_date (YYYY-MM-DD), quantity, available_shipping_date,
        item_condition_type, origin_type, origin_country_id,
        item_status, under18s_display, default_description
    }
    """
    wb = load_workbook(str(_TEMPLATE_PATH))
    ws = wb["Sheet1"]

    # 기본값 (미지정 시 합리적인 디폴트)
    cat = defaults.get("category_number", "")
    brand = defaults.get("brand_number", "")
    ship = defaults.get("shipping_number", "")
    end_date = defaults.get("end_date") or (date.today() + timedelta(days=365)).isoformat()
    qty = defaults.get("quantity", 100)
    asd = defaults.get("available_shipping_date", 7)
    condition = defaults.get("item_condition_type", "1")     # 1=새상품
    origin = defaults.get("origin_type", "2")                 # 2=해외
    country = defaults.get("origin_country_id", "KR")
    status = defaults.get("item_status", "Y")
    under18 = defaults.get("under18s_display", "N")
    default_desc = defaults.get("default_description", "")

    for i, r in enumerate(rows):
        target_row = DATA_START_ROW + i

        # 상품명 — 큐텐 fit title_jp 우선 (Phase 4-B LLM), 없으면 한국명/번역명 폴백
        item_name = (
            r.get("qoo10_title_jp")
            or r.get("product_name")
            or r.get("product_name_ko")
            or ""
        )
        _set(ws, target_row, "item_name", item_name)
        _set(ws, target_row, "price_yen", int(r.get("sell_price_jpy") or 0))
        _set(ws, target_row, "image_main_url", r.get("cover_image_url") or "")

        # 무게: 큐텐은 kg 단위 문자열. weight_g 가 +200g 적용된 등록용 무게
        weight_g = r.get("weight_g")
        if weight_g:
            weight_kg = round(float(weight_g) / 1000, 2)
            _set(ws, target_row, "item_weight", weight_kg)

        # 상세 — 마케팅 포인트 줄바꿈 포함 + notes/default_description
        marketing = r.get("qoo10_marketing") or []
        marketing_text = ""
        if isinstance(marketing, list) and marketing:
            marketing_text = "\n".join(f"・{m}" for m in marketing if m)
        notes = r.get("notes") or ""
        desc_parts = [s for s in [marketing_text, notes, default_desc] if s]
        desc = "\n\n".join(desc_parts)
        _set(ws, target_row, "item_description", desc)

        # 검색 키워드 — qoo10_tags 우선 (LLM 생성, 공백 join), 없으면 source 파싱
        tags = r.get("qoo10_tags") or []
        if isinstance(tags, list) and tags:
            sk = " ".join(t for t in tags if t)
        else:
            src = r.get("source") or ""
            sk = ""
            if ":" in src:
                sk = src.split(":", 1)[1].strip()
            elif r.get("search_keyword"):
                sk = r["search_keyword"]
        _set(ws, target_row, "search_keyword", sk)

        # 공통 기본값
        _set(ws, target_row, "category_number", cat)
        _set(ws, target_row, "brand_number", brand)
        _set(ws, target_row, "item_status", status)
        _set(ws, target_row, "end_date", end_date)
        _set(ws, target_row, "quantity", qty)
        _set(ws, target_row, "shipping_number", ship)
        _set(ws, target_row, "available_shipping_date", asd)
        _set(ws, target_row, "item_condition_type", condition)
        _set(ws, target_row, "origin_type", origin)
        if origin == "2":
            _set(ws, target_row, "origin_country_id", country)
        _set(ws, target_row, "under18s_display", under18)

    # 메모리로 저장
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue()
